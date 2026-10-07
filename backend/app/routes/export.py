import io
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import json

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER

from app.database import get_db
from app.models import ChatHistory, Document, User
from app.auth import get_current_user

router = APIRouter(prefix="/export", tags=["Export Tools"])

@router.get("/chat/{chat_id}/pdf")
def export_chat_to_pdf(
    chat_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Generates and downloads a clean, styled PDF containing the chat question,
    answer, and listed source documents.
    """
    chat = db.query(ChatHistory).filter(
        ChatHistory.id == chat_id,
        ChatHistory.user_id == current_user.id
    ).first()
    
    if not chat:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Chat history item not found"
        )
        
    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer, 
        pagesize=letter,
        rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    story = []
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle', 
        parent=styles['Heading1'], 
        alignment=TA_CENTER, 
        spaceAfter=20,
        textColor='#1E3A8A' # Dark blue
    )
    section_title_style = ParagraphStyle(
        'SectionTitle', 
        parent=styles['Heading2'], 
        spaceBefore=15, 
        spaceAfter=8,
        textColor='#2563EB' # Blue
    )
    body_style = ParagraphStyle(
        'Body', 
        parent=styles['Normal'], 
        fontSize=10, 
        leading=14, 
        spaceAfter=12
    )
    q_style = ParagraphStyle(
        'Question',
        parent=body_style,
        fontName='Helvetica-Bold',
        backColor='#F3F4F6', # Light gray bg
        borderPadding=10,
        spaceAfter=15
    )
    
    # PDF Content
    story.append(Paragraph("Smart Document Intelligence Platform", title_style))
    story.append(Paragraph("AI Answer Export", ParagraphStyle('Sub', parent=styles['Normal'], alignment=TA_CENTER, spaceAfter=20)))
    
    story.append(Paragraph("User Query:", section_title_style))
    story.append(Paragraph(chat.question, q_style))
    
    story.append(Paragraph("AI Response:", section_title_style))
    # Replace newlines with break tags for PDF styling
    formatted_answer = chat.answer.replace("\n", "<br/>")
    story.append(Paragraph(formatted_answer, body_style))
    
    # Parse sources
    if chat.sources:
        try:
            sources_list = json.loads(chat.sources)
            if sources_list:
                story.append(Spacer(1, 15))
                story.append(Paragraph("Sources Referenced:", section_title_style))
                for s in sources_list:
                    src_text = f"• <b>{s.get('filename', 'Unknown')}</b> (Page {s.get('page_number', '?')}) - Similarity: {s.get('score', 0.0):.2f}"
                    story.append(Paragraph(src_text, body_style))
        except Exception:
            pass
            
    doc.build(story)
    pdf_buffer.seek(0)
    
    filename = f"chat_export_{chat_id}.pdf"
    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/summary/{doc_id}/docx")
def export_summary_to_docx(
    doc_id: int,
    summary_data: dict, # Expects short_summary, detailed_summary, bullet_summary in body
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Generates and downloads a Word Document (.docx) containing the summary elements.
    """
    document = db.query(Document).filter(
        Document.id == doc_id,
        Document.user_id == current_user.id
    ).first()
    
    if not document:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found"
        )
        
    docx_doc = DocxDocument()
    
    # Title
    docx_doc.add_heading('Smart Document Summary', level=0)
    docx_doc.add_paragraph(f"Document Name: {document.filename}")
    docx_doc.add_paragraph(f"Date Created: {document.upload_date.strftime('%Y-%m-%d %H:%M:%S')}")
    docx_doc.add_paragraph().add_run().bold = True
    
    # Short summary
    docx_doc.add_heading('Executive Summary', level=1)
    docx_doc.add_paragraph(summary_data.get("short_summary", "No executive summary provided."))
    
    # Detailed Summary
    docx_doc.add_heading('Detailed Analysis', level=1)
    docx_doc.add_paragraph(summary_data.get("detailed_summary", "No detailed summary provided."))
    
    # Bullet Summary
    docx_doc.add_heading('Key Takeaways', level=1)
    bullets = summary_data.get("bullet_summary", "")
    for line in bullets.split('\n'):
        line = line.strip()
        if not line:
            continue
        # Remove markdown bullet syntax if present
        if line.startswith('* ') or line.startswith('- ') or line.startswith('• '):
            line = line[2:]
        docx_doc.add_paragraph(line, style='List Bullet')
        
    docx_buffer = io.BytesIO()
    docx_doc.save(docx_buffer)
    docx_buffer.seek(0)
    
    filename = f"summary_export_{doc_id}.docx"
    return StreamingResponse(
        docx_buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
